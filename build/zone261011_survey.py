"""전수조사 1차 판정(구역 안 537지번 A~E, 2026-10-13 구역계 3안 방 작성)을 지번 좌표에 붙여 지도 층으로 넣는다.

Usage: python3 build/zone261011_survey.py JUDGE_XLSX LEDGER_XLSX
  JUDGE_XLSX  = 김해_초광역/시청보고준비_261015/01_전수조사/김해원도심_전수조사_1차판정표_261013.xlsx  (「1차 판정표(지번)」 시트)
  LEDGER_XLSX = 김해_초광역/구역내건물주소_전기사용량_261013/김해원도심_구역내_건물주소_대장상세_내부작업용_261013.xlsx  (지번별 위도·경도)
책상 자료 1차 분류일 뿐 공실 확정이 아니다(최종은 층·호 단위 증빙 2종 이상). 공개 페이지라 점포 이름·소유·가격·사용량은 싣지 않는다.
순서: zone261011_extras.py 다음.
"""
import sys, json, re
from pathlib import Path
from collections import Counter
import openpyxl
from shapely.geometry import shape, Point

ROOT = Path(__file__).resolve().parents[1]
judge_xlsx, ledger_xlsx = (Path(a).expanduser() for a in sys.argv[1:3])
ZP = ROOT / 'data/gimhae-zone-261011.json'
Z = json.loads(ZP.read_text()); poly = shape(Z['boundary']['geometry'])

lw = openpyxl.load_workbook(ledger_xlsx, read_only=True, data_only=True)
rows = list(lw['건물별 대장 상세'].iter_rows(values_only=True)); h = rows[3]
coord = {}
for r in rows[4:]:
    d = dict(zip(h, r))
    if d.get('지번') and d.get('위도') and (d['법정동'], d['지번']) not in coord:
        coord[(d['법정동'], d['지번'])] = (round(float(d['경도']), 6), round(float(d['위도']), 6))

jw = openpyxl.load_workbook(judge_xlsx, read_only=True, data_only=True)
rows = list(jw['1차 판정표(지번)'].iter_rows(values_only=True)); h = rows[3]
K_CLOSED, K_GONE, K_NOW = '폐업 신호\n(인허가, 폐업일/업종)', '상가정보에서 사라짐\n(2026-03→06)', '현재 등록 점포 수\n(상가정보 2026-06)'
pts, missing = [], []
for r in rows[4:]:
    d = dict(zip(h, r))
    if not d.get('지번') or not d.get('1차 판정'): continue
    key = (d['법정동'], f"{d['법정동']} {d['지번']}")
    ll = coord.get(key)
    if not ll: missing.append(key[1]); continue
    cls = str(d['1차 판정'])[0]
    closed, gone, now, rank = (d.get(k) for k in (K_CLOSED, K_GONE, K_NOW, '시 13곳 순위'))
    sig = [x for x in ['폐업 신호 %s' % closed if closed else '', '상가정보에서 사라짐' if gone else '',
                       '등록 점포 %s곳' % now if now else '', '시 우선순위 %s번' % rank if rank else ''] if x]
    pts.append({'type': 'Feature', 'geometry': {'type': 'Point', 'coordinates': list(ll)},
                'properties': {'c': cls, 'jibun': key[1], 'road': d.get('도로명 주소') or '', 'use': d.get('주용도') or '', 'kind': d.get('용도 구분') or '',
                               'judge': str(d['1차 판정']), 'signals': ' · '.join(sig)}})
counts = Counter(p['properties']['c'] for p in pts)
assert len(pts) + len(missing) == 537, (len(pts), len(missing))
assert sum(poly.covers(Point(*p['geometry']['coordinates'])) for p in pts) >= len(pts) - 5   # 분할·합병 지번은 옆 지번 좌표(대장 엑셀 비고)
blob = json.dumps(pts, ensure_ascii=False)
assert not re.search(r'소유|의향|보증금|만 ?원|억|kWh|남실장', blob), 'public wording'
Z['survey'] = {'source': '전수조사 1차 판정표(2026-10-13) — 구역 안 지번 537곳을 건축물대장·인허가 폐업·상가정보·공개 매물·시 13곳으로 나눈 책상 분류',
               'rule': '장기 공실 = 2026.1.1 이전부터 공실, 객관 증빙 3종(수도·전기 저사용 / 영업자료 미발생 / 1년 단위 반복 미사용) 중 2종 이상 — 층·호 단위(국토부 RE:CORE 가이드라인 설명회본 2026.9.11)',
               'counts': {k: counts.get(k, 0) for k in 'ABCDE'}, 'missing': missing,
               'labels': {'A': '장기공실 유력', 'B': '확인 필요', 'C': '최근 공실(2026년)', 'D': '영업 중', 'E': '주거만'},
               'points': {'type': 'FeatureCollection', 'features': pts}}
ZP.write_text(json.dumps(Z, ensure_ascii=False, separators=(',', ':')))
print(json.dumps({'points': len(pts), 'missing': missing, 'counts': Z['survey']['counts']}, ensure_ascii=False))
